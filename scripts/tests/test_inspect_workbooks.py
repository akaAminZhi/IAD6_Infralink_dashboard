from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from scripts.etl import inspect_workbooks


def test_header_helpers_and_expected_field_matching() -> None:
    headers = inspect_workbooks.make_unique_headers(
        ("Unique ID", "", "Unique ID"),
        4,
    )

    assert headers == ["Unique ID", "Unnamed: 2", "Unique ID_2", "Unnamed: 4"]
    assert inspect_workbooks.last_non_blank_index((None, "value", None)) == 1
    assert inspect_workbooks.last_non_blank_index((None, "")) == -1
    assert inspect_workbooks.make_json_value(datetime(2026, 7, 30, 8, 15)) == (
        "2026-07-30T08:15:00"
    )

    matches = inspect_workbooks.match_expected_fields(
        ["Unique ID", "NETA Complete: Completed", "Missing"],
        {"Unique ID", " neta   complete: completed ", "Other"},
    )
    assert matches == [
        {"expected": "Unique ID", "actual": "Unique ID", "match_type": "exact"},
        {
            "expected": "NETA Complete: Completed",
            "actual": " neta   complete: completed ",
            "match_type": "normalized",
        },
        {"expected": "Missing", "actual": None, "match_type": None},
    ]


def test_inspect_sheet_finds_header_counts_columns_and_samples() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "EXPORT"
    sheet.append([None, None, None])
    sheet.append(["Unique ID", "Status", "Status"])
    for index in range(12):
        sheet.append([f"IAD06-EQ-{index + 1}", "Active", index])

    result = inspect_workbooks.inspect_sheet(sheet)

    assert result["sheet_name"] == "EXPORT"
    assert result["header_row_number"] == 2
    assert result["row_count"] == 12
    assert result["column_count"] == 3
    assert result["column_headers"] == ["Unique ID", "Status", "Status_2"]
    assert len(result["first_10_rows"]) == 10
    assert result["first_10_rows"][0] == {
        "Unique ID": "IAD06-EQ-1",
        "Status": "Active",
        "Status_2": 0,
    }


def test_inspect_empty_sheet() -> None:
    workbook = Workbook()
    result = inspect_workbooks.inspect_sheet(workbook.active)

    assert result["row_count"] == 0
    assert result["column_count"] == 0
    assert result["column_headers"] == []
    assert result["first_10_rows"] == []


def test_inspect_workbook_and_main_write_output(
    monkeypatch,
    tmp_path: Path,
) -> None:
    raw_dir = tmp_path / "raw_data"
    raw_dir.mkdir()
    paths: dict[str, str] = {}

    for key in inspect_workbooks.WORKBOOK_KEYS:
        path = raw_dir / f"{key}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "EXPORT"
        sheet.append(["Unique ID", "Status"])
        sheet.append(["IAD06-EQ-1", "Active"])
        workbook.save(path)
        paths[key] = str(path)

    output_path = tmp_path / "frontend" / "public" / "data" / "inspection.json"
    monkeypatch.setattr(inspect_workbooks, "PROJECT_ROOT", tmp_path)
    monkeypatch.setattr(inspect_workbooks, "OUTPUT_PATH", output_path)
    monkeypatch.setattr(inspect_workbooks, "get_input_files", lambda: paths)

    workbook_result = inspect_workbooks.inspect_workbook(
        "system_elements",
        paths["system_elements"],
    )
    assert workbook_result["sheet_names"] == ["EXPORT"]
    assert workbook_result["sheets"][0]["row_count"] == 1
    assert "Unique ID" in workbook_result["expected_fields_present"]
    assert "NETA Test Report" in workbook_result["expected_fields_missing"]

    inspect_workbooks.main()
    payload = json.loads(output_path.read_text(encoding="utf-8"))
    assert [item["workbook_key"] for item in payload["workbooks"]] == (
        inspect_workbooks.WORKBOOK_KEYS
    )
    assert payload["project_context"]["dashboard_orientation"] == "PDM-centric"


def test_inspect_workbook_rejects_missing_file(tmp_path: Path) -> None:
    missing = tmp_path / "missing.xlsx"
    try:
        inspect_workbooks.inspect_workbook("cases", str(missing))
    except FileNotFoundError as exc:
        assert str(missing) in str(exc)
    else:
        raise AssertionError("Expected FileNotFoundError")
