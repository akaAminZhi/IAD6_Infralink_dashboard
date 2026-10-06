"""Attach confirmed skid energisation flags and preserve source/matching diagnostics."""

from __future__ import annotations

from collections import defaultdict
import os
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

try:
    from .json_utils import file_metadata
    from .normalize_modules import normalize_header
except ImportError:
    from json_utils import file_metadata
    from normalize_modules import normalize_header


PROJECT_ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_NAME = "Offsite Energisation SKID tracker.xlsx"
HEADERS = {
    "pdm_name": "Skid Name",
    "energised_offsite": "Energised Offsite? (Yes/No)",
    "energised_onsite": "Onsite Energisation Performed? (Yes/No)",
}
STATUS_FIELDS = ("energised_offsite", "energised_onsite")


def tracker_workbook_path() -> Path:
    root = Path(
        os.environ.get(
            "IAD6_EPS_TRACKER_ROOT", PROJECT_ROOT.parent / "IAD6_EPS_Testing_Tracker"
        )
    )
    return root / "Excel" / WORKBOOK_NAME


def pdm_key(value: Any) -> str:
    return " ".join(str(value or "").upper().split())


def cds_alias(value: Any) -> str:
    # The shipping tracker also uses e.g. E6-120-03-CDS-R for PRIMARY-CDS-R.
    # Restrict this alias to complete numbered CDS skid names, retaining -R.
    return re.sub(
        r"^(IAD06-PDM-[A-Z]+\d+-\d{3}-\d{2})-PRIMARY-CDS(-R)?$",
        r"\1-CDS\2",
        pdm_key(value),
    )


def read_records(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    found_header = False
    try:
        for sheet in workbook:
            columns: dict[str, int] | None = None
            for row_number, row in enumerate(
                sheet.iter_rows(values_only=True), start=1
            ):
                if columns is None:
                    header = {
                        normalize_header(str(value)): i
                        for i, value in enumerate(row)
                        if value is not None
                    }
                    if all(
                        normalize_header(value) in header for value in HEADERS.values()
                    ):
                        columns = {
                            key: header[normalize_header(value)]
                            for key, value in HEADERS.items()
                        }
                        found_header = True
                    elif row_number >= 30:
                        break
                    continue
                raw = {key: row[index] for key, index in columns.items()}
                if all(
                    value is None or str(value).strip() == "" for value in raw.values()
                ):
                    continue
                record: dict[str, Any] = {
                    "sheet": sheet.title,
                    "row": row_number,
                    "pdm_name": str(raw["pdm_name"] or "").strip(),
                    "raw_values": {
                        key: None if value is None else str(value)
                        for key, value in raw.items()
                    },
                    "invalid_fields": [],
                }
                for field in STATUS_FIELDS:
                    value = (
                        str(raw[field] if raw[field] is not None else "")
                        .strip()
                        .casefold()
                    )
                    record[field] = {"yes": True, "no": False}.get(value)
                    if value not in {"", "yes", "no"}:
                        record["invalid_fields"].append(field)
                records.append(record)
    finally:
        workbook.close()
    if not found_header:
        raise ValueError(f"Missing skid energisation headers in {path}")
    return records


def attach_energisation(
    pdms: list[dict[str, Any]],
    path: Path | None = None,
) -> dict[str, Any]:
    source = path if path is not None else tracker_workbook_path()
    for pdm in pdms:
        for field in STATUS_FIELDS:
            pdm[field] = None
    if not source.is_file():
        return {
            "available": False,
            "source_file": None,
            "source_path": str(source),
            "records": [],
        }

    records = read_records(source)
    exact: dict[str, list[int]] = defaultdict(list)
    aliases: dict[str, list[int]] = defaultdict(list)
    for index, pdm in enumerate(pdms):
        exact[pdm_key(pdm.get("pdm_name"))].append(index)
        aliases[cds_alias(pdm.get("pdm_name"))].append(index)
    assignments: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        key = pdm_key(record["pdm_name"])
        candidates = (
            (exact.get(key, []) or aliases.get(cds_alias(key), [])) if key else []
        )
        record["candidate_pdms"] = [pdms[index]["pdm_name"] for index in candidates]
        record["match_status"] = "unmatched" if not candidates else "ambiguous"
        if len(candidates) == 1:
            assignments[candidates[0]].append(record)
            record["match_status"] = "matched"
            record["match_method"] = "exact" if exact.get(key) else "cds_alias"
    for index, matches in assignments.items():
        if len(matches) != 1:
            for record in matches:
                record["match_status"] = "ambiguous"
            continue
        for field in STATUS_FIELDS:
            pdms[index][field] = matches[0][field]
    return {"available": True, "source_file": file_metadata(source), "records": records}
